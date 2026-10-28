"""Safe local ingestion for text, Markdown, PDF and DOCX uploads."""
from __future__ import annotations

import json
from pathlib import Path
import re
from typing import Optional
from threading import RLock
from uuid import uuid4

from fastapi import HTTPException, UploadFile

from .corpus import AcademicDocument
from .schemas import UserRole
from .ollama import OllamaClient
from .vision_ingestion import process_syllabus_with_vision

SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx"}
MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MANIFEST_FILENAME = "knowledge-base.json"
KNOWLEDGE_BASE_MANAGERS = frozenset({UserRole.administrator, UserRole.staff, UserRole.faculty})
_manifest_lock = RLock()


def can_manage_knowledge_base(role: UserRole) -> bool:
    """Only institutional staff may add a source to the shared knowledge base."""
    return role in KNOWLEDGE_BASE_MANAGERS


def _manifest_path(upload_dir: Path) -> Path:
    return upload_dir / MANIFEST_FILENAME


def _manifest_records(upload_dir: Path) -> list[dict]:
    path = _manifest_path(upload_dir)
    if not path.exists():
        return []
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []
    return payload if isinstance(payload, list) else []


def persist_uploaded_document(document: AcademicDocument, upload_dir: Path) -> None:
    """Persist extracted text and metadata so an upload survives an app restart.

    The original file remains on disk too. Persisting the extracted text avoids
    depending on optional PDF/DOCX parsers during a later startup.
    """
    upload_dir.mkdir(parents=True, exist_ok=True)
    record = {
        "id": document.id,
        "title": document.title,
        "source": document.source,
        "text": document.text,
        "allowed_roles": [role.value for role in document.allowed_roles],
    }
    with _manifest_lock:
        records = [item for item in _manifest_records(upload_dir) if item.get("id") != document.id]
        records.append(record)
        path = _manifest_path(upload_dir)
        temporary_path = path.with_suffix(".tmp")
        temporary_path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary_path.replace(path)


def load_persisted_documents(upload_dir: Path, client: Optional[OllamaClient] = None) -> list[AcademicDocument]:
    """Restore valid persisted uploads, syncing with files on disk so uploads survive any restart or reload."""
    upload_dir.mkdir(parents=True, exist_ok=True)
    restored: list[AcademicDocument] = []
    seen_sources: set[str] = set()
    seen_texts: set[str] = set()
    records_to_persist: list[dict] = []
    manifest_updated = False

    with _manifest_lock:
        records = _manifest_records(upload_dir)
        for record in records:
            if not isinstance(record, dict):
                continue
            identifier = record.get("id")
            title = record.get("title")
            source = record.get("source")
            text = record.get("text")
            raw_roles = record.get("allowed_roles")
            if not (isinstance(identifier, str) and identifier.strip() and
                    isinstance(title, str) and title.strip() and
                    isinstance(source, str) and source.strip() and
                    isinstance(text, str) and text.strip()):
                continue
            roles: list[UserRole] = []
            for raw_role in raw_roles if isinstance(raw_roles, list) else []:
                try:
                    roles.append(UserRole(raw_role))
                except ValueError:
                    continue
            if not roles:
                roles = list(UserRole)
            if text in seen_texts:
                continue
            seen_texts.add(text)
            seen_sources.add(source)
            doc = AcademicDocument(
                id=identifier,
                title=title,
                source=source,
                text=text,
                allowed_roles=tuple(roles),
            )
            restored.append(doc)
            records_to_persist.append(record)

        # Scan upload_dir for any supported files on disk that may not be in the manifest yet
        for disk_file in sorted(upload_dir.iterdir()):
            if disk_file.name.startswith(".") or disk_file.name == MANIFEST_FILENAME:
                continue
            if disk_file.suffix.lower() not in SUPPORTED_SUFFIXES:
                continue
            source_tag = f"upload:{disk_file.name}"
            if source_tag in seen_sources:
                continue
            try:
                suffix = disk_file.suffix.lower()
                if suffix == ".pdf" and client:
                    extracted = process_syllabus_with_vision(disk_file, client)
                else:
                    raw_bytes = disk_file.read_bytes()
                    if not raw_bytes or len(raw_bytes) > MAX_UPLOAD_BYTES:
                        continue
                    extracted = _extract_text(raw_bytes, suffix)

                extracted_text = re.sub(r"\n{3,}", "\n\n", extracted).strip() if extracted else ""
                if not extracted_text or len(extracted_text) < 25 or extracted_text in seen_texts:
                    continue

                clean_stem = re.sub(r"-[a-f0-9]{8,}$", "", disk_file.stem)
                clean_stem = re.sub(r"[^a-zA-Z0-9_-]+", "-", clean_stem).strip("-") or "document"
                title = clean_stem.replace("-", " ").replace("_", " ").title()

                doc_id = f"upload_{uuid4().hex[:12]}"
                new_doc = AcademicDocument(
                    id=doc_id,
                    title=title,
                    source=source_tag,
                    text=extracted_text,
                    allowed_roles=tuple(UserRole),
                )
                seen_sources.add(source_tag)
                seen_texts.add(extracted_text)
                restored.append(new_doc)
                records_to_persist.append({
                    "id": new_doc.id,
                    "title": new_doc.title,
                    "source": new_doc.source,
                    "text": extracted_text,
                    "allowed_roles": [r.value for r in new_doc.allowed_roles],
                })
                manifest_updated = True
            except Exception:
                continue

        if manifest_updated or (records and len(records_to_persist) != len(records)):
            path = _manifest_path(upload_dir)
            temporary_path = path.with_suffix(".tmp")
            temporary_path.write_text(json.dumps(records_to_persist, ensure_ascii=False, indent=2), encoding="utf-8")
            temporary_path.replace(path)

    return restored


def _extract_text(data: bytes, suffix: str) -> str:
    if suffix in {".txt", ".md"}:
        return data.decode("utf-8", errors="replace")
    if suffix == ".pdf":
        # Fallback to pypdf if vision is not used
        try:
            from pypdf import PdfReader
            from io import BytesIO
            return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)
        except ImportError as error:
            raise HTTPException(501, "PDF support needs the documents extra: pip install -e '.[documents]'") from error
    if suffix == ".docx":
        try:
            from docx import Document
            from io import BytesIO
            return "\n".join(paragraph.text for paragraph in Document(BytesIO(data)).paragraphs)
        except ImportError as error:
            raise HTTPException(501, "DOCX support needs the documents extra: pip install -e '.[documents]'") from error
    raise HTTPException(415, "Supported formats are .md, .txt, .pdf, and .docx.")


async def ingest_upload(upload: UploadFile, upload_dir: Path, roles: list[UserRole], client: Optional[OllamaClient] = None) -> AcademicDocument:
    suffix = Path(upload.filename or "document.txt").suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise HTTPException(415, "Supported formats are .md, .txt, .pdf, and .docx.")
    data = await upload.read()
    if not data:
        raise HTTPException(400, "The upload is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Uploads are limited to 8 MB in this starter deployment.")

    upload_dir.mkdir(parents=True, exist_ok=True)
    safe_stem = re.sub(r"[^a-zA-Z0-9_-]+", "-", Path(upload.filename or "document").stem).strip("-") or "document"
    stored_name = f"{safe_stem}-{uuid4().hex[:8]}{suffix}"
    stored_path = upload_dir / stored_name
    stored_path.write_bytes(data)

    if suffix == ".pdf" and client:
        extracted = process_syllabus_with_vision(stored_path, client)
    else:
        extracted = _extract_text(data, suffix)

    text = re.sub(r"\n{3,}", "\n\n", extracted).strip() if extracted else ""
    if len(text) < 25:
        raise HTTPException(422, "No usable text could be extracted from this document.")

    title = safe_stem.replace("-", " ").replace("_", " ").title()
    return AcademicDocument(
        id=f"upload_{uuid4().hex[:12]}",
        title=title,
        source=f"upload:{stored_name}",
        text=text,
        allowed_roles=tuple(roles or list(UserRole)),
    )
    
